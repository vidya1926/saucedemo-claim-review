from pathlib import Path
import re

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "journal-manuscript.md"
OUTPUT = ROOT / "Hallucination Research_paper.docx"


def make_box(ax, center, size, text, fill="#EAF2F8", edge="#2E5266", font_size=8):
    width, height = size
    x, y = center
    box = FancyBboxPatch(
        (x - width / 2, y - height / 2),
        width,
        height,
        boxstyle="round,pad=0.12,rounding_size=0.16",
        linewidth=1.0,
        edgecolor=edge,
        facecolor=fill,
    )
    ax.add_patch(box)
    ax.text(
        x, y, text, ha="center", va="center", fontsize=font_size,
        color="#172B3A", wrap=True, linespacing=1.2,
    )


def connect(ax, start, end, color="#526D82", curve=0):
    arrow = FancyArrowPatch(
        start,
        end,
        arrowstyle="-|>",
        mutation_scale=9,
        linewidth=1.0,
        color=color,
        connectionstyle=f"arc3,rad={curve}",
        shrinkA=2,
        shrinkB=2,
    )
    ax.add_patch(arrow)


def render_figures():
    figures = ROOT / "figures"
    figures.mkdir(exist_ok=True)

    fig, ax = plt.subplots(figsize=(3.2, 5.5))
    fig.subplots_adjust(left=0.035, right=0.965, top=0.985, bottom=0.025)
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 16)
    ax.axis("off")
    make_box(ax, (5, 14.7), (8.9, 1.25),
             "1. Scenario + equivalent prompts\nreviewed for the same conditions and outcome",
             "#EAF2F8", font_size=7.5)
    make_box(ax, (5, 12.5), (8.9, 1.25),
             "2. One fixed LLM configuration\nTest A + Test B; preserve raw outputs and settings",
             "#EAF2F8", font_size=7.5)
    make_box(ax, (5, 10.3), (8.9, 1.25),
             "3. Human claim mapping\nnormalize assertions to shared propositions",
             "#EAF2F8", font_size=7.5)
    nodes = [
        (1.75, "4. Consistency\ncompare\npaired claims"),
        (5.0, "5. Grounding\ncheck reviewed\nfacts"),
        (8.25, "6. Playwright\nrecord execution\nseparately"),
    ]
    for x, label in nodes:
        make_box(ax, (x, 7.7), (2.75, 1.65), label,
                 "#F4F7FA", font_size=6.1)
        connect(ax, (5, 9.65), (x, 8.55))
    make_box(ax, (5, 5.0), (8.9, 1.25),
             "7. Human review and adjudication\nsupported / contradicted / unresolved",
             "#FFF4D6", edge="#8A6D1D", font_size=7.4)
    for x, _ in nodes:
        connect(ax, (x, 6.85), (5, 5.65))
    make_box(ax, (5, 2.35), (8.9, 1.3),
             "8. Compare approaches and report\nmetrics, uncertainty, and limitations",
             "#E7F3E8", edge="#3D7044", font_size=7.4)
    connect(ax, (5, 4.35), (5, 3.05))
    path = figures / "research-workflow.png"
    fig.savefig(path, dpi=240, facecolor="white")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(3.2, 5.5))
    fig.subplots_adjust(left=0.035, right=0.965, top=0.985, bottom=0.025)
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 16)
    ax.axis("off")
    make_box(ax, (2.55, 14.55), (4.5, 1.25),
             "Paired prompts\n+ generated tests",
             "#EAF2F8", font_size=6.5)
    make_box(ax, (7.45, 14.55), (4.5, 1.25),
             "Independently reviewed\nevidence catalog",
             "#FFF4D6", edge="#8A6D1D", font_size=7.1)
    make_box(ax, (2.55, 12.15), (4.5, 1.4),
             "Manual claim mapping\nshared proposition\nidentifiers",
             "#EAF2F8", font_size=6.2)
    make_box(ax, (7.45, 12.15), (4.5, 1.4),
             "Evidence facts\nsource, date, context,\nartifact",
             "#FFF4D6", edge="#8A6D1D", font_size=6.2)
    connect(ax, (2.55, 13.9), (2.55, 12.85))
    connect(ax, (7.45, 13.9), (7.45, 12.85))

    branches = [
        (1.75, "Consistency\npaired\nclaims"),
        (5.0, "Grounding\ncompare with\nreviewed facts"),
        (8.25, "Playwright\nexecution\nobservations"),
    ]
    for x, label in branches:
        fill, edge = ("#F4F7FA", "#2E5266")
        if x == 5.0:
            fill, edge = "#FFF9E9", "#8A6D1D"
        elif x == 8.25:
            fill, edge = "#F4F7FA", "#617987"
        make_box(ax, (x, 9.55), (2.85, 1.7), label,
                 fill, edge=edge, font_size=6.0)
    connect(ax, (2.55, 11.45), (1.75, 10.4))
    connect(ax, (4.55, 11.45), (5.0, 10.4))
    connect(ax, (7.05, 11.45), (5.0, 10.4))
    connect(ax, (2.9, 11.45), (8.25, 10.4), curve=-0.1)

    make_box(ax, (5, 6.75), (8.9, 1.25),
             "Review queue + visualizer\nkeep signals distinguishable",
             "#EAF2F8", font_size=7.1)
    for x, _ in branches:
        connect(ax, (x, 8.65), (5, 7.4))
    make_box(ax, (5, 4.25), (8.9, 1.25),
             "Human adjudication\nsupported / contradicted /\nunresolved",
             "#FFF4D6", edge="#8A6D1D", font_size=6.8)
    connect(ax, (5, 6.1), (5, 4.9))
    make_box(ax, (5, 1.7), (8.9, 1.3),
             "Evaluation\nexecution-only | consistency-only\n"
             "grounding-only | combined",
             "#E7F3E8", edge="#3D7044", font_size=6.8)
    connect(ax, (5, 3.6), (5, 2.35))
    path = figures / "prototype-architecture.png"
    fig.savefig(path, dpi=240, facecolor="white")
    plt.close(fig)


def add_inline_runs(paragraph, text):
    text = text.translate(str.maketrans({
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u2026": "...",
        "\u2192": "->",
        "\u2260": "!=",
    }))
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")
    for part in pattern.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*"):
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def set_run_font(run, size=9.5):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)


def add_field(paragraph, instruction):
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    paragraph._p.append(field)


lines = SOURCE.read_text(encoding="utf-8").splitlines()
render_figures()
doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.68)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.68)
section.right_margin = Inches(0.68)
section.header_distance = Inches(0.3)
section.footer_distance = Inches(0.3)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal.font.size = Pt(9.2)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.line_spacing = 1.0
normal.paragraph_format.space_after = Pt(3)

for style_name, font_size in (("Heading 1", 10), ("Heading 2", 9.2)):
    style = doc.styles[style_name]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(font_size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(6)
    style.paragraph_format.space_after = Pt(2)
    style.paragraph_format.keep_with_next = True

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.CENTER
header_run = header.add_run("JOURNAL MANUSCRIPT DRAFT | PLAYWRIGHT UI TEST CLAIMS")
set_run_font(header_run, 7.5)
header_run.bold = True
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(footer.add_run("CSCI 6803 | October 2026 | "), 7.5)
add_field(footer, "PAGE")

title = next(line[2:].strip() for line in lines if line.startswith("# "))
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(6)
p.paragraph_format.space_after = Pt(7)
r = p.add_run(title)
r.bold = True
set_run_font(r, 15)

metadata = [line.strip() for line in lines if line.startswith("**Author:**")
            or line.startswith("**Affiliation:**")]
for line in metadata:
    value = re.sub(r"\*\*", "", line)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    add_inline_runs(p, value)
    for run in p.runs:
        set_run_font(run, 9)

for prefix in ("**Thesis progress:**", "**Status date:**"):
    line = next((line.strip() for line in lines if line.startswith(prefix)), None)
    if line:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        add_inline_runs(p, line.replace("**", ""))
        for run in p.runs:
            set_run_font(run, 8)

status = next((line for line in lines if line.startswith("**Manuscript status:**")), "")
if status:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    add_inline_runs(p, status.replace("**", ""))
    for run in p.runs:
        set_run_font(run, 8)

course = next((line for line in lines if line.startswith("**Course context:**")), "")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(7)
add_inline_runs(p, course.replace("**", ""))
for run in p.runs:
    set_run_font(run, 8)

abstract_heading = next(i for i, line in enumerate(lines) if line.strip() == "## Abstract")
abstract_text = next(line.strip() for line in lines[abstract_heading + 1:] if line.strip())
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
r = p.add_run("Abstract - ")
r.bold = True
add_inline_runs(p, abstract_text)

keywords = next((line for line in lines if line.startswith("**Index Terms—**")), None)
if keywords:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline_runs(p, keywords.replace("**", ""))
    for run in p.runs:
        set_run_font(run, 8.5)

# Start the article body in two columns, following the supplied IEEE-style example.
body_section = doc.add_section(WD_SECTION.CONTINUOUS)
body_section.top_margin = Inches(0.68)
body_section.bottom_margin = Inches(0.65)
body_section.left_margin = Inches(0.68)
body_section.right_margin = Inches(0.68)
body_section.header_distance = Inches(0.3)
body_section.footer_distance = Inches(0.3)
columns = OxmlElement("w:cols")
columns.set(qn("w:num"), "2")
columns.set(qn("w:space"), "300")
body_section._sectPr.append(columns)

inside_body = False
i = 0
while i < len(lines):
    raw = lines[i]
    stripped = raw.strip()
    if stripped.casefold() == "## introduction":
        inside_body = True
    if not inside_body:
        i += 1
        continue
    if not stripped or stripped == "---":
        i += 1
        continue
    if stripped.startswith("# "):
        i += 1
        continue
    if stripped.startswith("**Author:**") or stripped.startswith("**Affiliation:**") \
            or stripped.startswith("**Thesis progress:**") or stripped.startswith("**Status date:**") \
            or stripped.startswith("**Manuscript status:**") or stripped.startswith("**Course context:**") \
            or stripped.startswith("**Index Terms—**"):
        i += 1
        continue
    if stripped.startswith("```mermaid"):
        i += 1
        while i < len(lines) and not lines[i].strip().startswith("```"):
            i += 1
        i += 1
        continue
    image_match = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", stripped)
    if image_match:
        alt_text, image_path = image_match.groups()
        image = ROOT / image_path
        if not image.is_file():
            raise FileNotFoundError(f"Manuscript figure is missing: {image}")
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        p.add_run().add_picture(str(image), width=Inches(3.1))
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_after = Pt(5)
        caption.paragraph_format.keep_with_next = True
        add_inline_runs(caption, alt_text)
        for run in caption.runs:
            set_run_font(run, 7.8)
            run.italic = True
        i += 1
        continue
    if stripped.startswith("## "):
        heading_text = stripped[3:]
        p = doc.add_paragraph(style="Heading 1")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline_runs(p, heading_text)
        i += 1
        continue
    if stripped.startswith("### "):
        p = doc.add_paragraph(style="Heading 2")
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_inline_runs(p, stripped[4:])
        i += 1
        continue
    if stripped.startswith("- "):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.12)
        p.paragraph_format.first_line_indent = Inches(-0.1)
        p.paragraph_format.space_after = Pt(2)
        p.add_run("- ")
        add_inline_runs(p, stripped[2:])
        i += 1
        continue
    if re.match(r"^\d+\.\s", stripped):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(2)
        add_inline_runs(p, stripped)
        i += 1
        continue
    if re.match(r"^\[\d+\]\s", stripped):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.17)
        p.paragraph_format.first_line_indent = Inches(-0.17)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.0
        add_inline_runs(p, stripped)
        for run in p.runs:
            set_run_font(run, 8)
        i += 1
        continue

    chunks = [stripped]
    i += 1
    while i < len(lines):
        following = lines[i].strip()
        if not following or following.startswith(("# ", "## ", "### ", "- ")) \
                or re.match(r"^\[\d+\]\s", following):
            break
        chunks.append(following)
        i += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline_runs(p, " ".join(chunks))
    for run in p.runs:
        set_run_font(run, 9.2)

doc.core_properties.title = title
doc.core_properties.subject = "Journal-style research protocol and prototype manuscript"
doc.core_properties.author = ""
doc.save(OUTPUT)
print(f"Created {OUTPUT} ({OUTPUT.stat().st_size} bytes)")
