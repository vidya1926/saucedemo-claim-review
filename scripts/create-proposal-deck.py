from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "research-proposal.pptx"

SLIDE_W = 13.333
SLIDE_H = 7.5

NAVY = RGBColor(16, 35, 56)
INK = RGBColor(31, 48, 68)
BLUE = RGBColor(61, 107, 255)
TEAL = RGBColor(38, 183, 169)
GOLD = RGBColor(239, 177, 74)
CORAL = RGBColor(224, 101, 91)
PAPER = RGBColor(247, 249, 252)
WHITE = RGBColor(255, 255, 255)
MUTED = RGBColor(103, 119, 138)
LINE = RGBColor(222, 229, 237)
PALE_BLUE = RGBColor(232, 238, 255)
PALE_TEAL = RGBColor(226, 246, 243)
PALE_GOLD = RGBColor(251, 242, 221)
PALE_CORAL = RGBColor(251, 233, 231)


def inch(value):
    return Inches(value)


def set_background(slide, color=PAPER):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(slide, x, y, w, h, text, size=18, color=INK, bold=False,
             font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
             margin=0.04):
    box = slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = inch(margin)
    frame.margin_right = inch(margin)
    frame.margin_top = inch(margin)
    frame.margin_bottom = inch(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_rich_lines(slide, x, y, w, h, lines, size=16, color=INK,
                   bullet=False, gap=7, left_margin=0.12):
    box = slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = inch(left_margin)
    tf.margin_right = inch(0.12)
    tf.margin_top = inch(0.08)
    tf.margin_bottom = inch(0.06)
    for index, line in enumerate(lines):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        paragraph.text = f"•  {line}" if bullet else line
        paragraph.level = 0
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = color
        paragraph.space_after = Pt(gap)
    return box


def add_card(slide, x, y, w, h, title, body_lines, accent=BLUE,
             fill=WHITE, body_size=16, title_size=18):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, inch(x), inch(y), inch(w), inch(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = LINE
    shape.line.width = Pt(0.8)
    stripe = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, inch(x), inch(y + 0.08), inch(0.07), inch(h - 0.16)
    )
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = accent
    stripe.line.fill.background()
    add_text(slide, x + 0.22, y + 0.18, w - 0.42, 0.42, title,
             size=title_size, color=NAVY, bold=True)
    add_rich_lines(slide, x + 0.2, y + 0.72, w - 0.36, h - 0.84,
                   body_lines, size=body_size, color=INK, bullet=True)
    return shape


def add_header(slide, title, kicker=None):
    if kicker:
        add_text(slide, 0.62, 0.32, 11.9, 0.28, kicker.upper(),
                 size=10, color=TEAL, bold=True)
        title_y = 0.72
    else:
        title_y = 0.46
    add_text(slide, 0.62, title_y, 12.1, 0.68, title,
             size=28, color=NAVY, bold=True)
    rule = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, inch(0.64), inch(1.39), inch(0.8), inch(0.055)
    )
    rule.fill.solid()
    rule.fill.fore_color.rgb = TEAL
    rule.line.fill.background()


def add_footer(slide, slide_no, source=None):
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, inch(0.62), inch(7.08), inch(12.08), inch(0.012)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = LINE
    line.line.fill.background()
    if source:
        add_text(slide, 0.66, 7.12, 10.9, 0.2, source,
                 size=8, color=MUTED)
    add_text(slide, 11.95, 7.1, 0.65, 0.22, f"{slide_no:02d}",
             size=9, color=MUTED, align=PP_ALIGN.RIGHT)


def add_slide(prs, title, kicker=None, source=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, title, kicker)
    add_footer(slide, len(prs.slides), source)
    return slide


def add_pill(slide, x, y, w, label, fill, color=NAVY):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, inch(x), inch(y), inch(w), inch(0.36)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    add_text(slide, x + 0.06, y + 0.035, w - 0.12, 0.27, label,
             size=10, color=color, bold=True, align=PP_ALIGN.CENTER,
             valign=MSO_ANCHOR.MIDDLE, margin=0)
    return shape


def add_box(slide, x, y, w, h, title, subtitle, accent):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, inch(x), inch(y), inch(w), inch(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = accent
    shape.line.width = Pt(1.3)
    add_text(slide, x + 0.08, y + 0.12, w - 0.16, 0.39, title,
             size=14, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, x + 0.09, y + 0.53, w - 0.18, h - 0.59, subtitle,
             size=10, color=MUTED, align=PP_ALIGN.CENTER)
    return shape


def add_arrow(slide, x1, y1, x2, y2, color=MUTED, width=1.8):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, inch(x1), inch(y1), inch(x2), inch(y2)
    )
    connector.line.color.rgb = color
    connector.line.width = Pt(width)
    connector.line.end_arrowhead = True
    return connector


def build_deck():
    prs = Presentation()
    prs.slide_width = inch(SLIDE_W)
    prs.slide_height = inch(SLIDE_H)
    prs.core_properties.title = (
        "Detecting Hallucinations in AI-Generated Software Test Cases"
    )
    prs.core_properties.subject = (
        "Research proposal: metamorphic testing for Playwright UI tests"
    )
    prs.core_properties.author = "Research proposal"
    prs.core_properties.keywords = (
        "LLM, hallucination, UI testing, metamorphic testing, Playwright"
    )

    # 1. Title
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, NAVY)
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, inch(0), inch(0), inch(0.18), inch(SLIDE_H)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = TEAL
    accent.line.fill.background()
    add_pill(slide, 0.82, 0.7, 1.82, "RESEARCH PROPOSAL", PALE_TEAL, NAVY)
    add_text(
        slide, 0.82, 1.48, 10.9, 1.78,
        "Detecting Hallucinations in\nAI-Generated Software Test Cases",
        size=34, color=WHITE, bold=True,
    )
    add_text(
        slide, 0.86, 3.55, 10.8, 0.78,
        "An evaluation of metamorphic testing for Playwright UI tests",
        size=21, color=RGBColor(191, 211, 229),
    )
    add_text(
        slide, 0.86, 5.62, 10.5, 0.7,
        "Paired prompts  •  Application evidence  •  Human-reviewed evaluation",
        size=15, color=RGBColor(196, 217, 231),
    )
    add_text(slide, 11.75, 6.72, 0.85, 0.28, "01", size=10,
             color=RGBColor(196, 217, 231), align=PP_ALIGN.RIGHT)

    # 2. Motivation
    slide = add_slide(prs, "A plausible test can still assert the wrong behavior",
                      "The problem")
    add_card(slide, 0.72, 1.75, 3.82, 3.85, "Generated test",
             ["May compile and run",
              "May use convincing locators and assertions",
              "Can encode an expectation the app does not satisfy"],
             accent=BLUE, body_size=16)
    add_card(slide, 4.78, 1.75, 3.82, 3.85, "Why it matters",
             ["False confidence in a test suite",
              "Incorrect requirements can become automated checks",
              "Failures can waste time if the expected behavior is invented"],
             accent=GOLD, body_size=16)
    add_card(slide, 8.84, 1.75, 3.78, 3.85, "Research challenge",
             ["A passing test is not proof of correctness",
              "A failing test may be flaky or misconfigured",
              "Compare behavioral claims and verify them independently"],
             accent=TEAL, body_size=16)
    add_text(slide, 0.82, 6.0, 11.7, 0.48,
             "Focus: unsupported user-visible behavior claims—not code style or exact selector matching.",
             size=17, color=NAVY, bold=True)

    # 3. Prior work and gap
    slide = add_slide(prs, "The proposal connects three research areas",
                      "State of the art")
    add_card(slide, 0.72, 1.7, 3.84, 3.6, "MetaQA",
             ["Metamorphic prompt mutations for LLM response consistency",
              "Foundation for comparing outputs from related prompts"],
             accent=BLUE, body_size=15)
    add_card(slide, 4.76, 1.7, 3.84, 3.6, "LLM web-test generation",
             ["Prior work generates navigation and form-filling tests",
              "Emphasis is often on generation and coverage"],
             accent=TEAL, body_size=15)
    add_card(slide, 8.8, 1.7, 3.82, 3.6, "Generated-test reliability",
             ["Unit-test research studies compilation, runtime and context errors",
              "Different domain from UI behavior claims"],
             accent=GOLD, body_size=15)
    band = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, inch(0.85), inch(5.72), inch(11.6), inch(0.82)
    )
    band.fill.solid()
    band.fill.fore_color.rgb = PALE_BLUE
    band.line.fill.background()
    add_text(slide, 1.08, 5.9, 11.12, 0.42,
             "Proposed gap: empirically evaluate metamorphic checks + application grounding for generated UI-test claims.",
             size=16, color=NAVY, bold=True, align=PP_ALIGN.CENTER)

    # 4. Aim and questions
    slide = add_slide(prs, "Measure whether the signals help—and where they fail",
                      "Aim & research questions")
    add_text(slide, 0.82, 1.7, 11.65, 0.7,
             "Aim: evaluate whether meaning-preserving prompt mutations, combined with independently reviewed app evidence, surface unsupported claims in AI-generated Playwright tests.",
             size=19, color=NAVY, bold=True)
    questions = [
        ("RQ1", "How often do equivalent prompts produce conflicting behavioral claims?"),
        ("RQ2", "How accurately do consistency checks identify reviewed claim errors?"),
        ("RQ3", "Does grounding add value over consistency-only or execution-only checks?"),
        ("RQ4", "What false alarms and unresolved cases remain?"),
    ]
    y = 2.66
    for index, (label, question) in enumerate(questions):
        fill = [PALE_BLUE, PALE_TEAL, PALE_GOLD, PALE_CORAL][index]
        add_pill(slide, 0.9, y + 0.06, 0.85, label, fill)
        add_text(slide, 1.98, y, 10.4, 0.5, question, size=17, color=INK)
        y += 0.79
    add_text(slide, 0.9, 6.12, 11.5, 0.42,
             "This evaluates a fixed pretrained LLM; it does not train a new model.",
             size=15, color=MUTED, bold=True)

    # 5. Architecture
    slide = add_slide(prs, "Architecture: compare claims, ground them, then review",
                      "Proposed system")
    # Main generation path
    top_y = 1.78
    box_w = 2.18
    gap = 0.72
    xs = [0.78, 3.68, 6.58, 9.48]
    add_box(slide, xs[0], top_y, box_w, 1.03, "Requirement",
            "Scenario + expected outcome", BLUE)
    add_box(slide, xs[1], top_y, box_w, 1.03, "Prompt pair",
            "Original + reviewed paraphrase", BLUE)
    add_box(slide, xs[2], top_y, box_w, 1.03, "Same LLM",
            "Fixed model + settings", TEAL)
    add_box(slide, xs[3], top_y, box_w, 1.03, "Playwright tests",
            "Generated test A and B", TEAL)
    for i in range(3):
        add_arrow(slide, xs[i] + box_w + 0.06, top_y + 0.52,
                  xs[i + 1] - 0.08, top_y + 0.52)

    add_box(slide, 4.62, 3.18, 3.2, 0.92, "Claim extraction",
            "Normalize observable behavior, not code", BLUE)
    add_arrow(slide, 10.55, 2.82, 7.65, 3.16)
    add_arrow(slide, 3.48, 2.82, 5.14, 3.16)

    add_box(slide, 0.92, 4.72, 3.22, 1.06, "Consistency check",
            "Do paired claims disagree?", BLUE)
    add_box(slide, 5.06, 4.72, 3.22, 1.06, "Grounding check",
            "Compare to independent evidence", TEAL)
    add_box(slide, 9.2, 4.72, 3.22, 1.06, "Playwright evidence",
            "Observed state / traces; signal only", GOLD)
    add_arrow(slide, 5.24, 4.1, 3.0, 4.69)
    add_arrow(slide, 6.25, 4.1, 6.65, 4.69)
    add_arrow(slide, 7.2, 4.1, 10.72, 4.69)
    add_arrow(slide, 4.14, 5.24, 5.04, 5.24)
    add_arrow(slide, 8.29, 5.24, 9.18, 5.24)
    add_text(slide, 1.0, 6.12, 11.1, 0.45,
             "Evidence catalog → human-reviewed labels → compare baselines and report metrics",
             size=15, color=NAVY, bold=True, align=PP_ALIGN.CENTER)

    # 6. Metamorphic rule
    slide = add_slide(prs, "Compare the meaning—not the exact script",
                      "Metamorphic relation")
    add_card(slide, 0.82, 1.78, 5.45, 3.82, "Meaning-preserving mutation",
             ["Keep actor, preconditions, action, and outcome unchanged",
              "Review each prompt pair before generation",
              "Compare normalized user-visible behavioral claims"],
             accent=BLUE, body_size=16)
    add_card(slide, 7.05, 1.78, 5.45, 3.82, "What should not be required",
             ["Identical Playwright code",
              "Identical locators or page-object design",
              "Identical formatting or implementation details"],
             accent=TEAL, body_size=16)
    add_pill(slide, 5.97, 3.12, 0.9, "≠", PALE_GOLD, NAVY)
    add_text(slide, 1.02, 6.0, 11.3, 0.48,
             "A disagreement is a signal to inspect. Two agreeing outputs may still both be wrong.",
             size=17, color=NAVY, bold=True, align=PP_ALIGN.CENTER)

    # 7. Study protocol
    slide = add_slide(prs, "Pilot on a public demo site with changing UI state",
                      "Study protocol")
    add_card(slide, 0.78, 1.72, 3.78, 3.85, "Target: SauceDemo",
             ["Public interactive demo; no Salesforce access needed",
              "Login, cart, sort, checkout, logout scenarios",
              "Record observation date and artifacts"],
             accent=TEAL, body_size=15)
    add_card(slide, 4.78, 1.72, 3.78, 3.85, "Planned sample",
             ["5 scenario families",
              "8 repeated prompt pairs per scenario",
              "Up to 40 pairs / 80 generated tests",
              "Feasibility target—not a power claim"],
             accent=BLUE, body_size=15)
    add_card(slide, 8.78, 1.72, 3.78, 3.85, "Audit trail",
             ["Prompts, model/version, settings",
              "Generated tests + normalized claims",
              "Evidence source and browser artifacts",
              "Reviewer label and rationale"],
             accent=GOLD, body_size=15)
    add_text(slide, 0.9, 6.0, 11.4, 0.42,
             "Prototype smoke tests exercise login, add/remove cart, and price sorting; research evidence still requires independent review.",
             size=14, color=MUTED, bold=True, align=PP_ALIGN.CENTER)

    # 8. Signals and labels
    slide = add_slide(prs, "Keep the three signals separate",
                      "Interpretation")
    add_card(slide, 0.82, 1.72, 3.78, 3.8, "Pair disagreement",
             ["Equivalent prompts imply opposite claims",
              "Label: inconsistency signal",
              "Not proof of hallucination"],
             accent=BLUE, fill=PALE_BLUE, body_size=15)
    add_card(slide, 4.78, 1.72, 3.78, 3.8, "Evidence comparison",
             ["Matches / contradicts supplied fact",
              "Missing fact → unresolved by this evidence",
              "Evidence quality must be reviewed"],
             accent=TEAL, fill=PALE_TEAL, body_size=15)
    add_card(slide, 8.74, 1.72, 3.78, 3.8, "Browser execution",
             ["Pass/fail, UI state, trace, screenshot",
              "May reveal technical issues",
              "Not the truth oracle by itself"],
             accent=GOLD, fill=PALE_GOLD, body_size=15)
    add_text(slide, 0.9, 5.9, 11.6, 0.5,
             "Outcome: a human review queue—not an automatic verdict.",
             size=20, color=NAVY, bold=True, align=PP_ALIGN.CENTER)

    # 9. Evaluation
    slide = add_slide(prs, "Compare four approaches against reviewed labels",
                      "Evaluation")
    approaches = [
        ("01", "Execution-only", "Browser outcomes and captured observations"),
        ("02", "Consistency-only", "Disagreement across equivalent prompts"),
        ("03", "Grounding-only", "Claim compared with evidence catalog"),
        ("04", "Combined", "Consistency + grounding under a predeclared rule"),
    ]
    y = 1.72
    for index, (number, title, detail) in enumerate(approaches):
        add_pill(slide, 0.88, y + 0.05, 0.74, number,
                 [PALE_BLUE, PALE_TEAL, PALE_GOLD, PALE_CORAL][index])
        add_text(slide, 1.86, y, 2.72, 0.38, title,
                 size=17, color=NAVY, bold=True)
        add_text(slide, 4.65, y, 7.55, 0.42, detail,
                 size=15, color=INK)
        y += 0.83
    add_card(slide, 0.92, 5.18, 11.5, 1.16, "Report",
             ["Confusion matrix, precision, recall, F1, false positives, unresolved cases; claim- and pair-level where appropriate."],
             accent=TEAL, fill=WHITE, body_size=14, title_size=16)

    # 10. Prototype status and limitations
    slide = add_slide(prs, "What exists now—and what remains research work",
                      "Current status")
    add_card(slide, 0.8, 1.72, 5.72, 3.96, "Implemented prototype",
             ["Offline claim-review CLI with JSON scenario format",
              "Surfaces disagreement, evidence conflict, and missing facts",
              "6 analyzer unit tests pass",
              "SauceDemo Playwright smoke suite: 4 tests pass"],
             accent=TEAL, body_size=15)
    add_card(slide, 6.8, 1.72, 5.72, 3.96, "Not yet established",
             ["No LLM integration or automatic claim extraction",
              "Example evidence is illustrative, not ground truth",
              "No labeled corpus or measured precision/recall",
              "No claim of validated detector effectiveness"],
             accent=CORAL, body_size=15)
    add_text(slide, 0.94, 6.03, 11.4, 0.45,
             "Do not use real account credentials or production customer data.",
             size=16, color=CORAL, bold=True, align=PP_ALIGN.CENTER)

    # 11. Work plan and contribution
    slide = add_slide(prs, "Next steps: validate the evidence before scaling",
                      "Plan & contribution")
    weeks = [
        ("1–2", "Confirm scope + review literature"),
        ("3", "Pilot prompts, labels, and evidence"),
        ("4–5", "Collect paired generations"),
        ("6", "Review and adjudicate claims"),
        ("7–8", "Analyze, report, and revise"),
    ]
    x = 0.82
    widths = [2.2, 2.2, 2.2, 2.2, 2.2]
    for i, (week, label) in enumerate(weeks):
        add_pill(slide, x, 1.92, 0.82, f"W{week}",
                 [PALE_BLUE, PALE_TEAL, PALE_GOLD, PALE_CORAL, PALE_BLUE][i])
        add_text(slide, x, 2.43, widths[i], 0.8, label,
                 size=14, color=INK, bold=True)
        if i < len(weeks) - 1:
            add_arrow(slide, x + widths[i] - 0.14, 2.31,
                      x + widths[i] + 0.14, 2.31, color=LINE, width=1.2)
        x += 2.48
    band = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, inch(0.9), inch(4.0), inch(11.5), inch(1.5)
    )
    band.fill.solid()
    band.fill.fore_color.rgb = NAVY
    band.line.fill.background()
    add_text(slide, 1.22, 4.27, 10.85, 0.95,
             "Contribution to evaluate: whether metamorphic checks add useful signal beyond grounding and execution—and how often they mislead.",
             size=21, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
             valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, 0.96, 5.95, 11.4, 0.5,
             "Novelty remains a research question; complete a systematic literature review before claiming priority.",
             size=14, color=MUTED, bold=True, align=PP_ALIGN.CENTER)

    # 12. References
    slide = add_slide(prs, "Selected starting references",
                      "References")
    references = [
        "Yang, B. et al. (2025). Hallucination Detection in Large Language Models with Metamorphic Relations. Proceedings of the ACM on Software Engineering, 2(FSE), 425–445. doi:10.1145/3715735",
        "Le, N.-K. et al. (2025). Automated Web Application Testing: End-to-End Test Case Generation with Large Language Models and Screen Transition Graphs. Proceedings of JSAI 2025. doi:10.11517/pjsai.JSAI2025.0_3K4IS2a03",
        "Gu, S. et al. (2024). TestART: Improving LLM-based Unit Testing via Co-evolution of Automated Generation and Repair Iteration. arXiv:2408.03095.",
        "Yin, X. et al. (2025). Enhancing LLM's Ability to Generate More Repository-Aware Unit Tests Through Precise Contextual Information Injection. arXiv:2501.07425.",
    ]
    add_rich_lines(slide, 0.85, 1.72, 11.65, 3.7, references,
                   size=15, color=INK, bullet=True, gap=16, left_margin=0.08)
    add_text(slide, 0.92, 5.7, 11.3, 0.72,
             "These are representative works from an initial scan, not a completed systematic review. Verify publisher records and expand the search before making a novelty claim.",
             size=15, color=MUTED, bold=True)

    prs.save(OUTPUT)
    print(f"Created {OUTPUT}")
    print(f"Slides: {len(prs.slides)}")


if __name__ == "__main__":
    build_deck()
